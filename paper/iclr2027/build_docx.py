"""Produce a .docx of the paper laid out like the ICLR PDF.

Word cannot reproduce TeX's typesetting, and nothing here pretends otherwise:
line breaks, justification and hyphenation will differ, and the line-number
ruler has no Word equivalent. What is reproduced is everything a reader or a
submission form actually measures. Page size, margins, header and footer
positions, body font and leading, and heading sizes are taken from the
compiled PDF rather than guessed:

    page            8.5 x 11 in
    left/right      1.5 in each, giving the venue's 5.5 in measure
    running head    0.386 in from the top
    body top        1.171 in
    body            Times New Roman 10pt on 11pt
    sections        12pt small caps, subsections 10pt small caps

Two substitutions are forced by the format and are the only content changes:

  * Figures point at the PNGs rather than the PDFs, because Word will not
    render a PDF image. They come from the same generator and the same data.
  * Citations are resolved by citeproc rather than the venue's .bst, which
    Word cannot run. The style is author-year, as the venue's is.

    python build_docx.py            # build, then verify
    python build_docx.py --no-verify

PANDOC overrides the converter path.
"""

import argparse
import os
import re
import pathlib
import shutil
import subprocess
import sys

HERE = pathlib.Path(__file__).resolve().parent
WORK = HERE / "docx"
OUT = WORK / "main.docx"
REFERENCE = WORK / "reference.docx"
SOURCE = WORK / "main_docx.tex"

PANDOC = os.environ.get(
    "PANDOC",
    str(pathlib.Path(os.path.expanduser("~"))
        / "AppData/Local/Temp/pandocdir/pandoc-3.5/pandoc.exe"))

RUNNING_HEAD = "Under review as a conference paper at ICLR 2027"

# Inches, measured off out/main.pdf.
PAGE_W, PAGE_H = 8.5, 11.0
MARGIN_SIDE = 1.5
HEADER_FROM_TOP = 0.386
BODY_TOP = 1.171
BODY_BOTTOM_MARGIN = 1.0
FOOTER_FROM_BOTTOM = 0.45

BODY_PT = 10
LEADING_PT = 11
SECTION_PT = 12
SUBSECTION_PT = 10
TITLE_PT = 17.28
FONT = "Times New Roman"


NEWLABEL = re.compile(r"\\newlabel\{([^}]+)\}\{\{([^}]*)\}\{([^}]*)\}")


def label_numbers():
    """label -> the number LaTeX assigned it, read from out/main.aux.

    pandoc has no cross-reference resolver, so it renders \\ref{eq:ceiling} as
    the literal "[eq:ceiling]". The numbers exist already: the LaTeX build
    writes them to the aux file. Reading them there keeps the two documents
    saying the same thing without a second numbering scheme.
    """
    aux = HERE / "out" / "main.aux"
    if not aux.exists():
        print("  no out/main.aux; build the PDF first or references stay raw")
        return {}
    labels = {}
    for key, number, _page in NEWLABEL.findall(
            aux.read_text(encoding="utf-8", errors="replace")):
        if key not in labels and number:
            labels[key] = number
    return labels


def _brace_group(text, open_index):
    """Index just past the group that opens at `open_index`.

    A regex cannot do this. Captions here nest two deep, as in
    `$\\rho_{\\max}$`, and a pattern allowing one level of nesting silently
    skipped two of the three table captions while appearing to work.
    """
    depth = 0
    for i in range(open_index, len(text)):
        if text[i] == "{":
            depth += 1
        elif text[i] == "}":
            depth -= 1
            if depth == 0:
                return i + 1
    raise ValueError("unbalanced braces from index %d" % open_index)


def number_captions(text, labels):
    """Prefix each caption with the number LaTeX gave its float.

    The float environment adds "Figure N:" on its own; pandoc emits the
    caption body alone, so the two documents would number differently or not
    at all.
    """
    out = []
    count = 0
    pos = 0
    caption = "\\caption{"
    while True:
        start = text.find(caption, pos)
        if start < 0:
            out.append(text[pos:])
            break
        body_open = start + len(caption) - 1
        body_end = _brace_group(text, body_open)
        body = text[body_open + 1:body_end - 1]

        after = text[body_end:body_end + 200]
        match = re.match(r"\s*\\label\{([^}]+)\}", after)
        number = labels.get(match.group(1)) if match else None
        kind = "Table" if text.rfind("\\begin{table}", 0, start) > \
            text.rfind("\\begin{figure}", 0, start) else "Figure"

        out.append(text[pos:start])
        if number:
            out.append("\\caption{\\textbf{%s %s:} %s}" % (kind, number, body))
            count += 1
        else:
            out.append(text[start:body_end])
        pos = body_end
    return "".join(out), count


def prepare_source():
    """main.tex, with the three changes the converter needs.

    Figures point at the PNGs, cross-references are resolved to the numbers
    LaTeX assigned, and captions get the "Figure N:" prefix that the LaTeX
    float environment adds automatically and pandoc does not.
    """
    WORK.mkdir(exist_ok=True)
    text = (HERE / "main.tex").read_text(encoding="utf-8")

    figures = 0
    for stem in sorted(p.stem for p in (HERE / "figs").glob("*.pdf")):
        old = "{figs/%s.pdf}" % stem
        if old in text:
            text = text.replace(old, "{figs/%s.png}" % stem)
            figures += 1

    labels = label_numbers()

    text, captioned = number_captions(text, labels)

    resolved = 0
    for key, number in labels.items():
        for macro in ("\\ref{%s}" % key, "\\eqref{%s}" % key):
            if macro in text:
                repl = number if macro.startswith("\\ref") else "(%s)" % number
                text = text.replace(macro, repl)
                resolved += 1

    SOURCE.write_text(text, encoding="utf-8")
    print("source: %d figures to png, %d captions numbered, "
          "%d cross-references resolved" % (figures, captioned, resolved))
    return figures


def _pt(value):
    from docx.shared import Pt
    return Pt(value)


def _field(paragraph, instruction):
    """Insert a Word field, which python-docx has no API for.

    Used for the page number: a literal digit would be wrong on every page but
    the first.
    """
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, end):
        run._r.append(el)
    return run


def _plain_style(style, size_pt, small_caps=False, bold=False):
    """Force a style to black Times at a given size, theme settings included.

    Setting `style.font.name` alone is not enough. Word's built-in heading and
    title styles reference the document theme for both font and colour, and a
    theme reference wins over the style's own value: the first Word render came
    out in blue Calibri Light headings with a rule under the title, none of
    which is in this document's design. The theme attributes have to be removed
    from the style's rPr, not overwritten.
    """
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    style.font.name = FONT
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.italic = False
    style.font.small_caps = small_caps
    style.font.color.rgb = RGBColor(0, 0, 0)

    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        if fonts.get(qn("w:" + attr)) is not None:
            del fonts.attrib[qn("w:" + attr)]
    for attr in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + attr), FONT)

    # Built-in Title carries a bottom border; nothing in this format does.
    ppr = style.element.find(qn("w:pPr"))
    if ppr is not None:
        borders = ppr.find(qn("w:pBdr"))
        if borders is not None:
            ppr.remove(borders)


def build_reference():
    """A reference document carrying the page setup and the styles.

    pandoc copies section properties, headers, footers and styles from here,
    so this is where the format lives.
    """
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Inches, Pt

    doc = docx.Document()
    section = doc.sections[0]
    section.page_width = Inches(PAGE_W)
    section.page_height = Inches(PAGE_H)
    section.left_margin = Inches(MARGIN_SIDE)
    section.right_margin = Inches(MARGIN_SIDE)
    section.top_margin = Inches(BODY_TOP)
    section.bottom_margin = Inches(BODY_BOTTOM_MARGIN)
    section.header_distance = Inches(HEADER_FROM_TOP)
    section.footer_distance = Inches(FOOTER_FROM_BOTTOM)

    head = section.header.paragraphs[0]
    head.text = RUNNING_HEAD
    head.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in head.runs:
        run.font.name = FONT
        run.font.size = Pt(BODY_PT)

    foot = section.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(foot, " PAGE ")
    for run in foot.runs:
        run.font.name = FONT
        run.font.size = Pt(BODY_PT)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(LEADING_PT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.space_before = Pt(0)

    # The venue sets headings in small caps at 12pt and 10pt, not bold; the
    # run-in headings at level 4 are bold and not small caps.
    for name, size, caps, bold in (("Heading 1", SECTION_PT, True, False),
                                   ("Heading 2", SUBSECTION_PT, True, False),
                                   ("Heading 3", SUBSECTION_PT, True, False),
                                   ("Heading 4", BODY_PT, False, True)):
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        _plain_style(style, size, small_caps=caps, bold=bold)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    for name, size in (("Title", TITLE_PT), ("Subtitle", SECTION_PT)):
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        _plain_style(style, size)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for name in ("Caption", "Author", "Abstract", "Body Text"):
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        _plain_style(style, 9 if name == "Caption" else BODY_PT)

    REFERENCE.parent.mkdir(exist_ok=True)
    doc.save(str(REFERENCE))
    print("reference: %s" % REFERENCE.name)


def convert():
    if not pathlib.Path(PANDOC).exists():
        print("no pandoc at %s; set PANDOC" % PANDOC, file=sys.stderr)
        return 1
    cmd = [PANDOC, str(SOURCE),
           "--from", "latex",
           "--to", "docx",
           "--reference-doc", str(REFERENCE),
           "--citeproc",
           "--bibliography", str(HERE / "references.bib"),
           "--resource-path", str(HERE),
           "--number-sections",
           "--wrap", "none",
           "-o", str(OUT)]
    print("$ pandoc ... -o %s" % OUT.name)
    return subprocess.call(cmd, cwd=str(HERE))


ANON_NOTICE = "Anonymous authors\nPaper under double-blind review"

# pandoc does not run the style file, so three things it cannot know have to
# be repaired afterwards.
NUMBER_PREFIX = re.compile(r"^\d+(?:\.\d+)*\s*$")


def postprocess():
    """Repair what the converter cannot get right on its own.

    1. `\\maketitle` and `\\thanks` are defined in the style file, which pandoc
       parses only as far as line 78 before giving up. The remains reach the
       document as a paragraph reading "maketitle thanks aketitle".
    2. `--number-sections` numbers every level, so the run-in headings come
       out as "1.0.0.1 Research objective." The venue leaves them unnumbered.
    3. The author block is the one in the source. In the PDF the style file
       replaces it with the double-blind notice, and the two must agree or the
       .docx is the less anonymous of the pair.
    """
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = docx.Document(str(OUT))
    removed = renumbered = 0

    for para in list(doc.paragraphs):
        text = para.text.strip()

        if "maketitle" in text and "thanks" in text:
            para._element.getparent().remove(para._element)
            removed += 1
            continue

        if para.style.name == "Heading 4":
            for run in para.runs:
                if NUMBER_PREFIX.match(run.text or ""):
                    run.text = ""
                    renumbered += 1
                    break
            # pandoc separates the number from the text with a tab.
            for run in para.runs:
                if run.text.startswith("\t"):
                    run.text = run.text.lstrip("\t")
                    break

    # The author block is the first Normal paragraph after the title.
    for i, para in enumerate(doc.paragraphs[:6]):
        if "Anonymous Author" in para.text or "anon@example" in para.text:
            for run in para.runs[1:]:
                run._element.getparent().remove(run._element)
            para.runs[0].text = ANON_NOTICE
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.runs[0].font.size = Pt(BODY_PT)
            break

    # The abstract heading is plain text in the conversion; centre it and set
    # it in caps as the style file does.
    for para in doc.paragraphs[:8]:
        if para.text.strip().lower() == "abstract":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.text = "ABSTRACT"
                run.font.size = Pt(SECTION_PT)
            break

    # An inline image in a paragraph whose line spacing is EXACTLY 11pt is
    # clipped to 11pt tall. Word drew every figure as a hairline strip of axis
    # labels until this was fixed. Only the paragraphs holding a drawing are
    # changed; the body keeps its exact leading.
    from docx.oxml.ns import qn
    unclipped = 0
    for para in doc.paragraphs:
        if not para._element.findall(".//" + qn("w:drawing")):
            continue
        ppr = para._element.get_or_add_pPr()
        spacing = ppr.find(qn("w:spacing"))
        if spacing is None:
            spacing = ppr.makeelement(qn("w:spacing"), {})
            ppr.append(spacing)
        # python-docx's line_spacing setters clear the attributes but leave
        # the empty element behind, and the style's rule then still applies.
        # These have to be written out.
        spacing.set(qn("w:lineRule"), "auto")
        spacing.set(qn("w:line"), "240")
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        unclipped += 1

    doc.save(str(OUT))
    print("postprocess: removed %d stray paragraph(s), unnumbered %d run-in "
          "heading(s), unclipped %d figure(s), replaced the author block"
          % (removed, renumbered, unclipped))


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--no-verify", action="store_true")
    args = ap.parse_args(argv)

    prepare_source()
    build_reference()
    if convert():
        return 1
    postprocess()
    print("wrote %s (%.0f KB)" % (OUT, OUT.stat().st_size / 1024))
    if args.no_verify:
        return 0
    return subprocess.call([sys.executable, str(HERE / "verify_docx.py"),
                            str(OUT)], cwd=str(HERE))


if __name__ == "__main__":
    raise SystemExit(main())
